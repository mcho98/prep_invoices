#!/usr/bin/env python3
"""Turn the timesheet workbook into an invoice-import CSV (same columns as
sample_invoice_import.csv).

One invoice per client. Each day and service with hours becomes one line:
  Customer             = client name (first line of the invoice only)
  Item(Product/Service)= Homemaking / Personal Care / Respite
  ItemDescription      = the date (YYYY-MM-DD)
  ItemQuantity         = hours
  ItemRate             = the client's $/hr
  ItemAmount           = hours x rate, rounded to cents

Workbook layout: one date row near the top (first date cell starts the month,
one column per day), then one section per client starting at a "Client Name"
label with the name to its right, holding a "rate $/hr" cell and the rows
"Homemaking hours", "Personal Care Hours", "Respite hours".

Usage:
  python prep_invoices.py mock_data.xlsx            # writes output/mock_data.csv
  python prep_invoices.py mock_data.xlsx --start 20260810 --end 20260820
  python prep_invoices.py mock_data.xlsx -o invoices.csv --invoice-date 2026-08-31 --terms-days 30 --first-invoice-no 1001
"""
import argparse
import csv
import datetime as dt
import json
import os
import sys
from decimal import Decimal, ROUND_HALF_UP

import openpyxl

SERVICES = {
    "homemaking hours": "Homemaking",
    "personal care hours": "Personal Care",
    "respite hours": "Respite",
}
CLIENT_LABEL = "client name"
RATE_LABEL = "rate $/hr"
CENT = Decimal("0.01")
OUTPUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output")

# config/ folder beside the script (or beside the executable / the .app when packaged with PyInstaller)
APP_DIR = os.path.dirname(os.path.abspath(sys.executable if getattr(sys, "frozen", False) else __file__))
if sys.platform == "darwin" and getattr(sys, "frozen", False) and APP_DIR.endswith(".app/Contents/MacOS"):
    APP_DIR = os.path.dirname(os.path.dirname(os.path.dirname(APP_DIR)))
CONFIG_FILE = os.path.join(APP_DIR, "config", "prep_invoices.json")
CHART_SHEET = "Family# Client# Chart"

HEADER = ["*InvoiceNo", "*Customer", "*InvoiceDate", "*DueDate", "Terms", "Location", "Memo",
          "Item(Product/Service)", "ItemDescription", "ItemQuantity", "ItemRate", "*ItemAmount",
          "ItemTaxAmount"]


def norm(value):
    return value.strip().lower() if isinstance(value, str) else None


def as_date(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    return None


def to_int(value):
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return int(number) if number == int(number) else None


def read_chart(path):
    """{(family number, client number): English name} from the client chart workbook.

    The chart is the sheet 'Family# Client# Chart' (or the first sheet if there is none), with a
    header row containing 'Family #', 'Client #' and 'Name'. Continuation rows are blank.
    """
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[CHART_SHEET] if CHART_SHEET in wb.sheetnames else wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        cols = None
        for i, row in enumerate(rows):
            labels = [norm(v) if isinstance(v, str) else None for v in row]
            if "family #" in labels and "client #" in labels and "name" in labels:
                cols = [labels.index(k) for k in ("family #", "client #", "name")]
                break
            if i >= 10:
                break
        if cols is None:
            raise ValueError("no header row with 'Family #', 'Client #' and 'Name' in the first 10 rows")
        names = {}
        for row in rows:
            row = list(row) + [None] * (max(cols) + 1 - len(row))
            fam, cli = to_int(row[cols[0]]), to_int(row[cols[1]])
            name = row[cols[2]]
            name = str(name).strip() if name is not None else ""
            if fam is not None and cli is not None and name:
                names.setdefault((fam, cli), name)
    finally:
        wb.close()
    if not names:
        raise ValueError("no family/client rows found")
    return names


def load_chart_path():
    """Chart path saved in the config file, or None if it is unset or the file is gone."""
    try:
        with open(CONFIG_FILE, encoding="utf-8") as f:
            path = json.load(f).get("chart")
    except (OSError, ValueError, AttributeError):
        return None
    return path if isinstance(path, str) and os.path.isfile(path) else None


def save_chart_path(path):
    try:
        os.makedirs(os.path.dirname(CONFIG_FILE), exist_ok=True)
        with open(CONFIG_FILE, "w", encoding="utf-8") as f:
            json.dump({"chart": os.path.abspath(path)}, f, indent=2)
    except OSError as e:
        print(f"Could not save {CONFIG_FILE}: {e}", file=sys.stderr)


def ask_for_chart():
    """Ask for the chart with a file picker (or the terminal if there is no display). None = cancelled."""
    print("The client chart is not set or can't be found.", file=sys.stderr)
    title = "Choose the client chart (Family# Client# Chart)"
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        path = filedialog.askopenfilename(title=title, filetypes=[("Excel files", "*.xlsx *.xlsm"),
                                                                   ("All files", "*.*")])
        root.destroy()
    except Exception:
        try:
            path = input(f"{title}\nPath: ").strip().strip("\"'")
        except EOFError:
            path = ""
    return path or None


def get_chart_path(explicit=None):
    """--chart if given, else the saved path, else ask (and save the answer). Exits if none."""
    if explicit:
        if not os.path.isfile(explicit):
            sys.exit(f"Client chart not found: {explicit}")
        save_chart_path(explicit)
        return explicit
    path = load_chart_path()
    if not path:
        path = ask_for_chart()
        if not path or not os.path.isfile(path):
            sys.exit("A client chart is needed to give clients their English names.")
        save_chart_path(path)
    return path


def find_date_columns(ws):
    """[(column, date)] for each day of the month, starting from the first date cell."""
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20)):
        for cell in row:
            start = as_date(cell.value)
            if start is None:
                continue
            columns, day, col = [], start, cell.column
            while day.month == start.month:
                columns.append((col, day))
                col += 1
                day += dt.timedelta(days=1)
            return columns
    sys.exit("Could not find the date row (no date cell in the first 20 rows).")


def find_clients(ws):
    starts = []
    for row in ws.iter_rows():
        for cell in row:
            if norm(cell.value) == CLIENT_LABEL:
                name = ws.cell(cell.row, cell.column + 1).value
                # Family and client numbers sit two rows above the label: family one column left of it, client in it.
                key = (to_int(ws.cell(cell.row - 2, cell.column - 1).value) if cell.row > 2 and cell.column > 1 else None,
                       to_int(ws.cell(cell.row - 2, cell.column).value) if cell.row > 2 else None)
                starts.append((cell.row, str(name).strip() if name else f"(unnamed, row {cell.row})", key))
    return [(name, row, starts[i + 1][0] - 1 if i + 1 < len(starts) else ws.max_row, key)
            for i, (row, name, key) in enumerate(starts)]


def read_client(ws, name, first_row, last_row, date_columns, start=None, end=None):
    """Return (rate, [(date, service, hours)]). Exits on a missing rate or bad hours."""
    rate, service_rows = None, {}
    for row in ws.iter_rows(min_row=first_row, max_row=last_row):
        for cell in row:
            label = norm(cell.value)
            if label == RATE_LABEL and rate is None:
                rate = ws.cell(cell.row, cell.column + 1).value
            elif label in SERVICES:
                service_rows.setdefault(cell.row, SERVICES[label])
    lines = []
    for row, service in sorted(service_rows.items()):
        for col, day in date_columns:
            if (start and day < start) or (end and day > end):
                continue
            cell = ws.cell(row, col)
            if cell.value in (None, "", 0):
                continue
            if not isinstance(cell.value, (int, float)) or not 0 < cell.value <= 24:
                sys.exit(f"{name}: {cell.coordinate} has invalid hours ({cell.value!r}).")
            lines.append((day, service, cell.value))
    if lines and not (isinstance(rate, (int, float)) and rate > 0):
        sys.exit(f"{name}: no valid hourly rate next to '{RATE_LABEL}' (found {rate!r}).")
    lines.sort(key=lambda l: (l[0], l[1]))
    return rate, lines


def yyyymmdd(text):
    try:
        if len(text) != 8:  # strptime would accept short forms like 2026081
            raise ValueError
        return dt.datetime.strptime(text, "%Y%m%d").date()
    except ValueError:
        raise argparse.ArgumentTypeError(f"'{text}' is not a valid date; use YYYYMMDD, e.g. 20260815")


def fmt(n):
    """Number without trailing zeros (3.5 -> '3.5', 34.03 -> '34.03', 4.0 -> '4')."""
    return f"{n:f}".rstrip("0").rstrip(".") if isinstance(n, Decimal) else f"{n:g}"


def build_rows(ws, invoice_date, due_date, terms, first_no, start=None, end=None, chart=None):
    date_columns = find_date_columns(ws)
    rows, number = [], first_no
    for sheet_name, first, last, key in find_clients(ws):
        rate, lines = read_client(ws, sheet_name, first, last, date_columns, start, end)
        if not lines:
            print(f"Skipping {sheet_name}: no hours.", file=sys.stderr)
            continue
        name = sheet_name
        if chart is not None:
            if key in chart:
                name = chart[key]
            else:
                print(f"WARNING: family {key[0]} / client {key[1]} ({sheet_name}) is not in the client chart; "
                      f"using the timesheet name.", file=sys.stderr)
        total = Decimal(0)
        for i, (day, service, hours) in enumerate(lines):
            amount = (Decimal(str(hours)) * Decimal(str(rate))).quantize(CENT, rounding=ROUND_HALF_UP)
            total += amount
            head = [number, name, invoice_date.strftime("%d/%m/%Y"), due_date.strftime("%d/%m/%Y"), terms,
                    "", ""] if i == 0 else [number, "", "", "", "", "", ""]
            rows.append(head + [service, day.strftime("%Y-%m-%d"), fmt(hours), fmt(rate), f"{amount:.2f}", ""])
        print(f"Invoice {number}: {name}, {len(lines)} lines, ${total:,.2f}", file=sys.stderr)
        number += 1
    return rows


def write_invoices(workbook, output, invoice_date, terms_days, first_no, start=None, end=None, chart_path=None):
    """Read the workbook and write the invoice CSV. Returns the number of lines written.

    chart_path: client chart workbook used to replace timesheet names with the English chart names.
    """
    chart = None
    if chart_path:
        try:
            chart = read_chart(chart_path)
        except Exception as e:
            sys.exit(f"Could not read the client chart {chart_path}: {e}")
    ws = openpyxl.load_workbook(workbook, data_only=True).active
    due = invoice_date + dt.timedelta(days=terms_days)
    terms = f"Net {terms_days}" if terms_days else "Due on receipt"
    rows = build_rows(ws, invoice_date, due, terms, first_no, start, end, chart)
    os.makedirs(os.path.dirname(os.path.abspath(output)), exist_ok=True)
    with open(output, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(HEADER)
        w.writerows(rows)
    print(f"Wrote {len(rows)} lines to {output}", file=sys.stderr)
    return len(rows)


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("workbook")
    p.add_argument("-o", "--output",
                   help="CSV to write (default: output/<workbook name>.csv next to this script)")
    p.add_argument("--invoice-date", type=dt.date.fromisoformat, default=dt.date.today(),
                   help="YYYY-MM-DD (default: today)")
    p.add_argument("--start", type=yyyymmdd, help="only include days on or after this date (YYYYMMDD)")
    p.add_argument("--end", type=yyyymmdd, help="only include days on or before this date (YYYYMMDD)")
    p.add_argument("--terms-days", type=int, default=30, help="days until due (default 30; 0 = Due on receipt)")
    p.add_argument("--first-invoice-no", type=int, default=1001)
    p.add_argument("--chart", help="client chart workbook (.xlsx/.xlsm); saved to config/prep_invoices.json. "
                                   "Default: the saved path, asking for the file if it is missing")
    args = p.parse_args()
    if not args.output:
        stem = os.path.splitext(os.path.basename(args.workbook))[0]
        args.output = os.path.join(OUTPUT_DIR, stem + ".csv")

    if args.start and args.end and args.start > args.end:
        p.error("--start is after --end")
    write_invoices(args.workbook, args.output, args.invoice_date, args.terms_days,
                   args.first_invoice_no, args.start, args.end, get_chart_path(args.chart))


if __name__ == "__main__":
    main()
